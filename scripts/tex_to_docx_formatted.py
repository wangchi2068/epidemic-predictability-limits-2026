"""Regenerate 近临界疫情论文_Word版.docx replicating the reference format
(王驰123.docx): 宋体 body, first-line indent 420 twips (~2 chars), 9pt, headings.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

SRC = Path("paper_cn/main.tex")
OUT = Path("paper_cn/近临界疫情论文_Word版.docx")

tex = SRC.read_text(encoding="utf-8")
BS = chr(92)  # backslash


def math_to_unicode(s):
    """Convert a LaTeX math fragment to readable Unicode text."""
    b = BS
    s = s.strip()

    SYMBOLS = {
        "varepsilon": "\u03b5", "delta": "\u03b4", "Delta": "\u0394", "phi": "\u03c6",
        "Phi": "\u03a6", "gamma": "\u03b3", "Gamma": "\u0393", "sigma": "\u03c3",
        "Sigma": "\u03a3", "mu": "\u03bc", "ell": "l", "tau": "\u03c4",
        "eta": "\u03b7", "kappa": "\u03ba", "alpha": "\u03b1", "beta": "\u03b2",
        "lambda": "\u03bb", "theta": "\u03b8", "rho": "\u03c1", "omega": "\u03c9",
        "Omega": "\u03a9", "pi": "\u03c0", "partial": "d", "infty": "\u221e",
        "approx": "\u2248", "simeq": "\u2243", "sim": "~", "propto": "\u221d",
        "leq": "\u2264", "geq": "\u2265", "neq": "\u2260", "equiv": "\u2261",
        "gg": ">>", "ll": "<<", "times": "\u00d7", "cdot": "\u00b7",
        "pm": "\u00b1", "to": "\u2192", "rightarrow": "\u2192", "mapsto": "\u21a6",
        "mid": "|", "in": "\u2208", "notin": "\u2209", "subseteq": "\u2286",
        "forall": "\u2200", "exists": "\u2203", "sum": "\u2211", "prod": "\u220f",
        "int": "\u222b", "nabla": "\u2207", "dots": "\u2026", "ldots": "\u2026",
        "cdots": "\u2026", "circ": "\u00b7", "backslash": "\\", "emptyset": "\u2205",
        "log": "log", "ln": "ln", "exp": "exp", "max": "max", "min": "min",
        "sup": "sup", "inf": "inf", "lim": "lim", "lvert": "|", "rvert": "|",
        "quad": " ", "qquad": "  ", "neg": "\u00ac", "check": "\u02c7",
        "prime": "'", "ast": "*", "star": "*", "cdotp": "\u00b7",
        "dots": "\u2026", "colon": ":", "vert": "|", "shortmid": "|", "hat": "",
        "bar": "", "vec": "\u2192",
    }

    def read_group(text, i):
        depth = 1
        j = i + 1
        while j < len(text) and depth:
            if text[j] == "{":
                depth += 1
            elif text[j] == "}":
                depth -= 1
            j += 1
        return text[i + 1 : j - 1], j

    def conv(text):
        out = []
        i = 0
        n = len(text)
        while i < n:
            if text[i] == b and i + 1 < n:
                nxt = text[i + 1]
                if nxt in "(),[]{}":
                    out.append(nxt)
                    i += 2
                    continue
                if nxt in ",;:! ":
                    out.append(" ")
                    i += 2
                    continue
                if nxt.isalpha():
                    j = i + 1
                    while j < n and text[j].isalpha():
                        j += 1
                    cmd = text[i + 1 : j]
                    if j < n and text[j] == "*":
                        j += 1
                    args = []
                    k = j
                    while True:
                        while k < n and text[k] in " \t":
                            k += 1
                        if k < n and text[k] == "{":
                            content, k = read_group(text, k)
                            args.append(content)
                        else:
                            break
                    j = k
                    if cmd == "frac" and len(args) >= 2:
                        out.append("(" + conv(args[0]) + ")/(" + conv(args[1]) + ")")
                    elif cmd == "sqrt" and args:
                        out.append("\u221a(" + conv(args[0]) + ")")
                    elif cmd == "binom" and len(args) >= 2:
                        out.append("C(" + conv(args[0]) + ", " + conv(args[1]) + ")")
                    elif cmd in ("text", "mathrm", "operatorname", "mathbf", "boldsymbol"):
                        out.append(conv(args[0]) if args else "")
                    elif cmd == "mathbb":
                        out.append(args[0] if args else "")
                    elif cmd == "mathcal":
                        out.append(args[0] if args else "")
                    elif cmd in ("hat", "bar", "tilde"):
                        out.append(conv(args[0]) + {"hat": "\u02c6", "bar": "\u02c9", "tilde": "\u02dc"}[cmd] if args else "")
                    elif cmd == "overset" and len(args) >= 2:
                        out.append(conv(args[1]))
                    elif cmd == "underset" and len(args) >= 2:
                        out.append(conv(args[1]))
                    elif cmd in ("left", "right", "big", "Big", "Bigg", "Biggl", "Biggr"):
                        out.append("")
                    elif cmd in SYMBOLS:
                        out.append(SYMBOLS[cmd])
                    else:
                        pass  # drop unknown commands (label, begin, end, ...)
                    i = j
                    continue
                else:
                    out.append(nxt)
                    i += 2
                    continue
            out.append(text[i])
            i += 1
        return "".join(out)

    s = conv(s)
    s = s.replace("{", "").replace("}", "")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def eq_to_text(inner):
    """Convert a display-equation body to readable unicode text."""
    inner = re.sub(re.escape(BS) + r"label\{[^}]*\}", "", inner)
    for env in ("split", "align", "aligned", "gather"):
        inner = inner.replace(BS + "begin{" + env + "}", "").replace(BS + "end{" + env + "}", "")
    inner = inner.replace(BS + BS, " ; ").replace("&=", "= ").replace("& =", "= ")
    return math_to_unicode(inner)


def strip_cmd(s):
    """Convert a LaTeX line to readable text."""
    b = BS
    s = s.strip()
    s = re.sub(b + r"%.*", "", s)
    # command blocks: \cmd{...} -> delete (use escape on the backslash prefix)
    for cmd in (
        "citep",
        "citet",
        "label",
        "begin",
        "end",
        "caption",
        "section",
        "subsection",
        "subsubsection",
        "ref",
        "eqref",
        "input",
        "include",
        "bibliography",
        "bibliographystyle",
    ):
        s = re.sub(re.escape(b + cmd) + r"\{[^}]*\}", "", s)
    # display math
    s = re.sub(re.escape(b + "begin{equation*}"), " [公式] ", s)
    s = re.sub(re.escape(b + "end{equation*}"), "", s)
    s = re.sub(re.escape(b + "begin{align*}"), " [公式] ", s)
    s = re.sub(re.escape(b + "end{align*}"), "", s)
    # inline math: convert $...$ segments to readable unicode
    s = re.sub(r"\$([^$]*)\$", lambda m: math_to_unicode(m.group(1)), s)
    # simple commands without braces
    s = s.replace(b + "noindent", "")
    s = s.replace(b + "textbf{", "")
    s = s.replace(b + "emph{", "")
    s = s.replace(b + "item", "\u2022")
    s = re.sub(re.escape(b) + r"[a-zA-Z]+\*?", "", s)
    # leftover braces and punctuation
    s = s.replace("{", "").replace("}", "")
    s = s.replace("~", " ")
    s = (
        s.replace("[ht]", "")
        .replace("[htbp]", "")
        .replace("[h]", "")
        .replace("[t]", "")
    )
    s = s.replace("---", "\u2014").replace("--", "\u2013")
    s = s.replace("``", "\u201c").replace("''", "\u201d")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def set_song_font(run, size_pt=9, bold=False):
    run.font.name = "宋体"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "宋体")
    rfonts.set(qn("w:eastAsia"), "宋体")
    rfonts.set(qn("w:hAnsi"), "宋体")
    rfonts.set(qn("w:cs"), "Times New Roman")


def add_body_para(doc, text, indent_first=420, size=9, bold=False):
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    if indent_first:
        ind.set(qn("w:firstLine"), str(indent_first))
    ppr.append(ind)
    run = p.add_run(text)
    set_song_font(run, size_pt=size, bold=bold)
    return p


def add_heading_para(doc, text, size=12):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_song_font(run, size_pt=size, bold=True)
    return p


def heading_text(stripped):
    """Extract readable title from a section-style line."""
    inner = stripped.split("{", 1)[1].split("}", 1)[0]
    return strip_cmd(inner)


doc = Document()
equation_counter = [0]  # mutable counter for equation numbering
theorem_counter = [0]
lemma_counter = [0]

m_title = re.search(re.escape(BS + "title") + r"\{([^}]*)\}", tex)
if m_title:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(strip_cmd(m_title.group(1)))
    set_song_font(run, size_pt=15, bold=True)

body_start = tex.find(BS + "begin{document}")
body = tex[body_start:]

lines = body.split("\n")
i = 0
in_itemize = False
while i < len(lines):
    stripped = lines[i].strip()

    if stripped.startswith(BS + "begin{document}"):
        i += 1
        continue
    elif stripped.startswith(BS + "begin{lemma}"):
        lemma_counter[0] += 1
        add_heading_para(doc, f"引理{lemma_counter[0]}", size=10)
    elif stripped.startswith(BS + "begin{theorem5R}"):
        add_heading_para(doc, "定理5R", size=10)
    elif stripped.startswith(BS + "begin{theorem}"):
        theorem_counter[0] += 1
        add_heading_para(doc, f"定理{theorem_counter[0]}", size=10)
    elif stripped.startswith(BS + "begin{definition}"):
        add_heading_para(doc, "定义", size=10)
    elif stripped.startswith(BS + "begin{equation}"):
        if BS + "end{equation}" in stripped:
            inner = stripped[len(BS + "begin{equation}"):].split(
                BS + "end{equation}", 1
            )[0]
            eq_text = eq_to_text(inner)
            if eq_text:
                equation_counter[0] += 1
                n = equation_counter[0]
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(eq_text + "    (" + str(n) + ")")
                set_song_font(run, size_pt=9)
            i += 1
            continue
        # multi-line block: collect lines until the \end{equation}
        j = i + 1
        eq_lines = []
        while j < len(lines) and BS + "end{equation}" not in lines[j]:
            eq_lines.append(lines[j])
            j += 1
        eq_text = " ".join(eq_to_text(ln) for ln in eq_lines if ln.strip())
        if eq_text:
            equation_counter[0] += 1
            n = equation_counter[0]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(eq_text + "    (" + str(n) + ")")
            set_song_font(run, size_pt=9)
        i = j + 1
        continue
    elif stripped.startswith(BS + "begin{abstract}"):
        j = i + 1
        abs_lines = []
        while j < len(lines) and BS + "end{abstract}" not in lines[j]:
            abs_lines.append(lines[j])
            j += 1
        abstract = " ".join(strip_cmd(ln) for ln in abs_lines if ln.strip())
        add_heading_para(doc, "摘要", size=12)
        add_body_para(doc, abstract)
        i = j
    elif stripped.startswith(BS + "section{"):
        add_heading_para(doc, heading_text(stripped), size=12)
    elif stripped.startswith(BS + "subsection{"):
        add_heading_para(doc, heading_text(stripped), size=11)
    elif stripped.startswith(BS + "subsubsection{"):
        add_heading_para(doc, heading_text(stripped), size=10)
    elif stripped.startswith(BS + "begin{itemize}"):
        in_itemize = True
    elif stripped.startswith(BS + "end{itemize}"):
        in_itemize = False
    elif stripped.startswith(BS + "begin{table}"):
        j = i + 1
        tab_lines = []
        while j < len(lines) and BS + "end{table}" not in lines[j]:
            tab_lines.append(lines[j])
            j += 1
        cap = ""
        for tl in tab_lines:
            m = re.search(re.escape(BS + "caption") + r"\{([^}]*)\}", tl)
            if m:
                cap = m.group(1)
        if cap:
            add_body_para(doc, "表：" + strip_cmd(cap), indent_first=0)
        for tl in tab_lines:
            if any(mark in tl for mark in ("midrule", "toprule", "bottomrule")):
                continue
            if "&" in tl and BS in tl:
                row = tl.replace(BS + BS, "").strip()
                cells = [strip_cmd(c) for c in row.split("&")]
                add_body_para(doc, "  |  ".join(cells), indent_first=0)
        i = j
    elif stripped.startswith(BS + "begin{figure}"):
        j = i + 1
        cap = ""
        while j < len(lines) and BS + "end{figure}" not in lines[j]:
            m = re.search(re.escape(BS + "caption") + r"\{([^}]*)\}", lines[j])
            if m:
                cap = m.group(1)
            j += 1
        if cap:
            add_body_para(doc, "图：" + strip_cmd(cap), indent_first=0)
        i = j
    elif stripped.startswith(BS + "end{document}"):
        break
    else:
        parts = re.split(
            r"(\\begin\{equation\}.*?\\end\{equation\})", stripped, flags=re.S
        )
        if len(parts) > 1:
            for part in parts:
                if part.startswith(BS + "begin{equation}"):
                    inner = part[len(BS + "begin{equation}") :]
                    inner = inner[: -len(BS + "end{equation}")]
                    eq_text = eq_to_text(inner)
                    if eq_text:
                        equation_counter[0] += 1
                        n = equation_counter[0]
                        p = doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = p.add_run(eq_text + "    (" + str(n) + ")")
                        set_song_font(run, size_pt=9)
                else:
                    txt = strip_cmd(part)
                    if txt:
                        if in_itemize:
                            add_body_para(doc, txt, indent_first=0)
                        else:
                            add_body_para(doc, txt)
            i += 1
            continue
        txt = strip_cmd(stripped)
        if txt:
            if in_itemize:
                add_body_para(doc, txt, indent_first=0)
            else:
                add_body_para(doc, txt)
    i += 1

doc.save(str(OUT))
print("saved:", OUT)
