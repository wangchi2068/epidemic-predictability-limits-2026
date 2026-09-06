"""Regenerate 近临界疫情论文_Word版.docx from the latest paper_cn/main.tex.

Lightweight LaTeX-to-DOCX via python-docx (pandoc unavailable):
- strips LaTeX commands to readable text
- maps sections/subsections to Word headings
- keeps paragraphs, itemize lists, tables (plain), math as inline text
- title (from preamble) + abstract included
"""

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

SRC = Path("paper_cn/main.tex")
OUT = Path("paper_cn/近临界疫情论文_Word版.docx")

tex = SRC.read_text(encoding="utf-8")


def strip_cmd(s):
    """Convert a LaTeX line to readable text."""
    s = s.strip()
    s = re.sub(r"\\%.*", "", s)
    s = re.sub(r"\\citep\{[^}]*\}", "[ref]", s)
    s = re.sub(r"\\citet\{[^}]*\}", "[ref]", s)
    s = re.sub(r"\\\((.*?)\\\)", r"\1", s)
    s = re.sub(r"\$(.*?)\$", r"\1", s)
    s = re.sub(r"\\begin\{equation\*\?\}", " [eq] ", s)
    s = re.sub(r"\\end\{equation\*\?\}", "", s)
    s = re.sub(r"\\begin\{align\*\?\}", " [eq] ", s)
    s = re.sub(r"\\end\{align\*\?\}", "", s)
    s = s.replace(r"\noindent", "")
    s = s.replace(r"\textbf{", "").replace("}", "")
    s = s.replace(r"\emph{", "").replace("}", "")
    s = s.replace(r"\item", "•")
    s = re.sub(r"\\[a-zA-Z]+\*?", "", s)
    s = s.replace("{", "").replace("}", "")
    s = s.replace("~", " ")
    s = s.replace("---", "—").replace("--", "–")
    s = s.replace("``", "\u201c").replace("''", "\u201d")
    s = re.sub(r"\s+", " ", s).strip()
    return s


doc = Document()
try:
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
except (KeyError, AttributeError):
    pass

# title from preamble
m_title = re.search(r"\\title\{([^}]*)\}", tex)
if m_title:
    doc.add_heading(strip_cmd(m_title.group(1)), level=0)

body_start = tex.find(r"\begin{document}")
body = tex[body_start:]

lines = body.split("\n")
i = 0
in_itemize = False
while i < len(lines):
    stripped = lines[i].strip()

    if stripped.startswith(r"\begin{abstract}"):
        j = i + 1
        abs_lines = []
        while j < len(lines) and r"\end{abstract}" not in lines[j]:
            abs_lines.append(lines[j])
            j += 1
        abstract = " ".join(strip_cmd(ln) for ln in abs_lines if ln.strip())
        doc.add_heading("摘要", level=1)
        doc.add_paragraph(abstract)
        i = j
    elif stripped.startswith(r"\section{"):
        title = re.sub(r"\\section\{|\}", "", stripped)
        doc.add_heading(title, level=1)
    elif stripped.startswith(r"\subsection{"):
        title = re.sub(r"\\subsection\{|\}", "", stripped)
        doc.add_heading(title, level=2)
    elif stripped.startswith(r"\subsubsection{"):
        title = re.sub(r"\\subsubsection\{|\}", "", stripped)
        doc.add_heading(title, level=3)
    elif stripped.startswith(r"\begin{itemize}"):
        in_itemize = True
    elif stripped.startswith(r"\end{itemize}"):
        in_itemize = False
    elif stripped.startswith(r"\begin{table}"):
        j = i + 1
        tab_lines = []
        while j < len(lines) and r"\end{table}" not in lines[j]:
            tab_lines.append(lines[j])
            j += 1
        cap = ""
        for tl in tab_lines:
            m = re.search(r"\\caption\{([^}]*)\}", tl)
            if m:
                cap = m.group(1)
        if cap:
            doc.add_paragraph("表：" + strip_cmd(cap))
        for tl in tab_lines:
            if any(mark in tl for mark in ("\\midrule", "\\toprule", "\\bottomrule")):
                continue
            if "&" in tl and "\\" in tl:
                row = tl.replace("\\\\", "").strip()
                cells = [strip_cmd(c) for c in row.split("&")]
                doc.add_paragraph(" | ".join(cells))
        i = j
    elif stripped.startswith(r"\begin{figure}"):
        j = i + 1
        cap = ""
        while j < len(lines) and r"\end{figure}" not in lines[j]:
            m = re.search(r"\\caption\{([^}]*)\}", lines[j])
            if m:
                cap = m.group(1)
            j += 1
        if cap:
            doc.add_paragraph("图：" + strip_cmd(cap))
        i = j
    elif stripped.startswith(r"\end{document}"):
        break
    else:
        txt = strip_cmd(stripped)
        if txt:
            if in_itemize:
                doc.add_paragraph(txt, style="List Bullet")
            else:
                doc.add_paragraph(txt)
    i += 1

doc.save(str(OUT))
print("saved:", OUT)
